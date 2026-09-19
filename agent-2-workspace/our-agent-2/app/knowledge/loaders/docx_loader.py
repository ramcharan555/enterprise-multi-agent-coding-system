from pathlib import Path

from docx import Document as DocxDocument

from app.knowledge.loaders.base import DocumentLoader
from app.knowledge.models import Document, DocumentType


class DOCXDocumentLoader(DocumentLoader):
    """Loader for Microsoft Word DOCX documents."""

    SUPPORTED_EXTENSIONS = {".docx"}

    def load(self, path: Path) -> Document:
        if path.suffix.lower() not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type: {path.suffix}"
            )

        doc = DocxDocument(str(path))

        paragraphs = [
            paragraph.text
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        ]

        content = "\n\n".join(paragraphs)

        return Document(
            document_id=str(path.resolve()),
            title=path.stem,
            content=content,
            source=str(path),
            document_type=self._detect_document_type(path),
            metadata={
                "paragraph_count": len(paragraphs),
            },
        )

    @staticmethod
    def _detect_document_type(path: Path) -> DocumentType:
        name = path.name.lower()

        if "architecture" in name:
            return DocumentType.ARCHITECTURE

        if "adr" in name:
            return DocumentType.ADR

        if "api" in name:
            return DocumentType.API_DOCUMENTATION

        if "guide" in name or "guideline" in name:
            return DocumentType.DEVELOPER_GUIDELINE

        return DocumentType.GENERAL