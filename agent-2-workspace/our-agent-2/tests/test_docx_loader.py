from pathlib import Path

from docx import Document as DocxDocument

from app.knowledge.loaders import DOCXDocumentLoader
from app.knowledge.models import DocumentType


def create_test_docx(path: Path) -> None:
    document = DocxDocument()

    document.add_heading("API Architecture Guide", level=1)
    document.add_paragraph(
        "New APIs must use versioning."
    )
    document.add_paragraph(
        "Authentication must use OAuth2."
    )

    document.save(str(path))


def test_docx_loader(tmp_path: Path):
    docx_file = tmp_path / "architecture.docx"
    create_test_docx(docx_file)

    loader = DOCXDocumentLoader()
    document = loader.load(docx_file)

    assert document.title == "architecture"
    assert document.document_type == DocumentType.ARCHITECTURE
    assert "New APIs must use versioning." in document.content
    assert "Authentication must use OAuth2." in document.content
    assert document.metadata["paragraph_count"] == 3


def test_docx_loader_rejects_non_docx(tmp_path: Path):
    text_file = tmp_path / "test.txt"
    text_file.write_text("hello", encoding="utf-8")

    loader = DOCXDocumentLoader()

    try:
        loader.load(text_file)
        assert False, "Expected ValueError"
    except ValueError as exc:
        assert "Unsupported file type" in str(exc)